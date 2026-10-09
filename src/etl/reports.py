"""Report generator shared by the API (on-demand) and the scheduler (weekly). No web-framework types in this module.

``build_report`` computes a report and returns bytes (CSV or multi-sheet Excel); ``record_run`` adds the report_runs row;
``generate_report`` does both and commits (the entry point for scripts and Airflow). Aggregates come from
``app.services.analytics`` (the same functions the dashboard uses), so report totals always equal the analytics endpoints.
The encounter-level report uses an explicit column allow-list: race and gender can never be exported.
"""
from __future__ import annotations

import io
import logging
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import Engine, and_, func, select
from sqlalchemy.orm import Session, aliased

from app.models import EncounterDiagnosis, ReportRun
from app.schemas.analytics import DENOMINATOR
from app.services import analytics
from app.services.views import V

log = logging.getLogger(__name__)
REPORT_TYPES = ("admissions_trend", "readmission_summary", "length_of_stay_summary", "encounters")
FORMATS = ("csv", "xlsx")
MAX_ENCOUNTER_EXPORT_ROWS = 50_000
CONTENT_TYPES = {"csv": "text/csv", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
SIMULATED_NOTE = "Admission and discharge dates are SIMULATED (the source dataset has no dates); trends are illustrative, not clinical findings."
ENCOUNTER_EXPORT_COLUMNS = [
    "encounter_id", "patient_nbr", "admission_date", "discharge_date", "age_group", "admission_type", "admission_source",
    "discharge_disposition", "medical_specialty", "payer_code", "time_in_hospital", "num_lab_procedures", "num_procedures",
    "num_medications", "number_outpatient", "number_emergency", "number_inpatient", "number_diagnoses", "max_glu_serum",
    "a1c_result", "med_changed", "diabetes_med", "readmitted", "readmitted_30d", "any_readmission",
    "is_readmission_eligible", "diag_1", "diag_2", "diag_3",
]
FORBIDDEN_COLUMNS = frozenset({"race", "gender"})
BREAKDOWNS = {"readmission_summary": ("age_group", "admission_type", "admission_source", "discharge_disposition",
                                      "specialty", "month"),
              "length_of_stay_summary": ("age_group", "admission_type", "admission_source", "specialty")}
HISTOGRAM_COLUMNS = [f"los_{d}_days" for d in range(1, 15)]


class ReportTooLarge(Exception):
    """The encounter-level export would exceed the row cap."""

    def __init__(self, rows: int, cap: int) -> None:
        self.rows, self.cap = rows, cap
        super().__init__(f"This export would contain {rows:,} encounters; the limit is {cap:,}. "
                         "Narrow it with date_from/date_to, age_group or admission_type_id.")


@dataclass(frozen=True)
class ReportResult:
    """A generated file held in memory."""

    content: bytes
    filename: str
    content_type: str
    row_count: int
    report_type: str
    fmt: str
    filters: dict[str, Any]


def _clean_filters(filters: dict[str, Any]) -> dict[str, Any]:
    return {k: (v.isoformat() if isinstance(v, date) else v) for k, v in filters.items() if v is not None}


def _group_frame(rows: list[dict[str, Any]], breakdown: str | None = None) -> pd.DataFrame:
    cols = ["label", "encounters", "eligible_encounters", "readmitted_30d", "readmission_rate", "small_n", "avg_length_of_stay"]
    df = pd.DataFrame(rows, columns=["id", *cols])[cols]
    if breakdown:
        df.insert(0, "breakdown", breakdown)
    return df


def _los_frame(rows: list[dict[str, Any]], breakdown: str) -> pd.DataFrame:
    base = pd.DataFrame([{k: r[k] for k in ("label", "encounters", "avg_length_of_stay", "min_length_of_stay",
                                            "max_length_of_stay", "small_n")} for r in rows])
    hist = pd.DataFrame([r["histogram"] for r in rows], columns=HISTOGRAM_COLUMNS)
    out = pd.concat([base, hist], axis=1)
    out.insert(0, "breakdown", breakdown)
    return out


def _encounter_frame(session: Session, filters: dict[str, Any]) -> pd.DataFrame:
    conditions = analytics.filter_conditions(filters)
    total = session.scalar(select(func.count()).select_from(V).where(*conditions)) or 0
    if total > MAX_ENCOUNTER_EXPORT_ROWS:
        raise ReportTooLarge(total, MAX_ENCOUNTER_EXPORT_ROWS)
    diag = [aliased(EncounterDiagnosis) for _ in range(3)]
    base_cols = [V.c[name] for name in ENCOUNTER_EXPORT_COLUMNS if not name.startswith("diag_")]
    stmt = select(*base_cols, *[d.icd9_code.label(f"diag_{i}") for i, d in enumerate(diag, start=1)]).select_from(V)
    for i, d in enumerate(diag, start=1):
        stmt = stmt.outerjoin(d, and_(d.encounter_id == V.c.encounter_id, d.position == i))
    if conditions:
        stmt = stmt.where(*conditions)
    rows = session.execute(stmt.order_by(V.c.admission_date, V.c.encounter_id)).all()
    df = pd.DataFrame([dict(r._mapping) for r in rows], columns=ENCOUNTER_EXPORT_COLUMNS)
    for col in ("med_changed", "diabetes_med", "readmitted_30d", "any_readmission", "is_readmission_eligible"):
        df[col] = df[col].astype(int)
    assert not FORBIDDEN_COLUMNS & set(df.columns)
    return df


def _summary_frame(summary: dict[str, Any]) -> pd.DataFrame:
    labels = {"total_encounters": "Total encounters", "unique_patients": "Unique patients",
              "avg_length_of_stay": "Average length of stay (days)", "avg_num_medications": "Average medications per encounter",
              "eligible_encounters": "Eligible encounters (readmission denominator)", "readmitted_30d": "Readmitted within 30 days (eligible)",
              "readmission_rate_30d": "30-day readmission rate (eligible)", "any_readmission_rate": "Any readmission rate (eligible)"}
    return pd.DataFrame({"measure": [labels[k] for k in labels], "value": [summary[k] for k in labels]})


def _tables(session: Session, report_type: str, filters: dict[str, Any]) -> tuple[dict[str, pd.DataFrame], pd.DataFrame, pd.DataFrame]:
    """Return (sheets for Excel, the single CSV table, the Summary sheet)."""
    summary = _summary_frame(analytics.summary(session, filters))
    if report_type == "admissions_trend":
        monthly = _group_frame(analytics.admissions_trend(session, "month", filters)).rename(columns={"label": "month"})
        yearly = _group_frame(analytics.admissions_trend(session, "year", filters)).rename(columns={"label": "year"})
        csv_table = pd.concat([monthly.rename(columns={"month": "period"}).assign(granularity="month"),
                               yearly.rename(columns={"year": "period"}).assign(granularity="year")])
        csv_table = csv_table[["granularity", "period", *[c for c in monthly.columns if c != "month"]]]
        return {"Monthly": monthly, "Yearly": yearly}, csv_table.reset_index(drop=True), summary
    if report_type == "readmission_summary":
        sheets = {b: _group_frame(analytics.readmissions(session, b, filters)) for b in BREAKDOWNS[report_type]}
        csv_table = pd.concat([_group_frame(analytics.readmissions(session, b, filters), b) for b in BREAKDOWNS[report_type]])
        return sheets, csv_table.reset_index(drop=True), summary
    if report_type == "length_of_stay_summary":
        raw = {b: analytics.length_of_stay(session, b, filters) for b in BREAKDOWNS[report_type]}
        sheets = {b: _los_frame(rows, b).drop(columns="breakdown") for b, rows in raw.items()}
        csv_table = pd.concat([_los_frame(rows, b) for b, rows in raw.items()])
        return sheets, csv_table.reset_index(drop=True), summary
    frame = _encounter_frame(session, filters)
    summary = pd.DataFrame({"measure": ["Encounters in this export"], "value": [len(frame)]})
    return {"Encounters": frame}, frame, summary


def _metadata(report_type: str, filters: dict[str, Any], row_count: int, triggered_by: str, now: datetime) -> pd.DataFrame:
    items = [("generated_at", now.isoformat(timespec="seconds")), ("report", report_type),
             ("filters", ", ".join(f"{k}={v}" for k, v in _clean_filters(filters).items()) or "none"),
             ("row_count", row_count), ("triggered_by", triggered_by), ("dates", SIMULATED_NOTE),
             ("denominator", DENOMINATOR),
             ("definition: 30-day readmission rate", "encounters with readmitted_30d among eligible encounters / eligible encounters"),
             ("definition: eligible encounter", "discharge disposition is not expired or hospice (ids 11, 13, 14, 19, 20, 21)"),
             ("definition: small_n", "true when a group has fewer than 11 encounters; do not over-interpret"),
             ("privacy", "race and gender are never included; encounter-level files contain patient_nbr for authorised roles only"),
             ("source", "UCI Diabetes 130-US Hospitals 1999-2008 (de-identified research data, educational use)")]
    return pd.DataFrame(items, columns=["item", "value"])


def _style_workbook(writer: pd.ExcelWriter, frames: dict[str, pd.DataFrame]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for name, df in frames.items():
        ws = writer.sheets[name]
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for idx, col in enumerate(df.columns, start=1):
            letter = get_column_letter(idx)
            longest = max([len(str(col))] + [len(str(v)) for v in df[col].head(500)])
            ws.column_dimensions[letter].width = min(max(12, longest + 2), 60)
            fmt = "yyyy-mm-dd" if "date" in str(col) else "0.00%" if "rate" in str(col) else None
            if fmt:
                for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                    row[0].number_format = fmt


def _to_excel(sheets: dict[str, pd.DataFrame], summary: pd.DataFrame, meta: pd.DataFrame) -> bytes:
    frames = {"Summary": summary, **sheets, "Metadata": meta}
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl", datetime_format="yyyy-mm-dd", date_format="yyyy-mm-dd") as writer:
        for name, df in frames.items():
            out = df.copy()
            for col in out.columns:
                if "date" in str(col) and len(out):
                    out[col] = pd.to_datetime(out[col])  # real Excel dates, not text
            out.to_excel(writer, sheet_name=name[:31], index=False)
        _style_workbook(writer, {n[:31]: f for n, f in frames.items()})
    return buf.getvalue()


def build_report(session: Session, report_type: str, filters: dict[str, Any], fmt: str,
                 triggered_by: str = "api", now: datetime | None = None) -> ReportResult:
    """Compute one report as CSV or Excel bytes. Raises ``ReportTooLarge`` for an oversized encounter export."""
    if report_type not in REPORT_TYPES:
        raise ValueError(f"unknown report type {report_type!r}")
    if fmt not in FORMATS:
        raise ValueError(f"unknown format {fmt!r}")
    now = now or datetime.now()
    sheets, csv_table, summary = _tables(session, report_type, filters)
    row_count = len(csv_table)
    if fmt == "csv":
        content = csv_table.to_csv(index=False).encode("utf-8")  # header row + data only: no metadata rows
    else:
        content = _to_excel(sheets, summary, _metadata(report_type, filters, row_count, triggered_by, now))
    filename = f"{report_type}_{now:%Y%m%d_%H%M%S}.{fmt}"
    log.info("built %s (%d rows, %d bytes) for %s", filename, row_count, len(content), triggered_by)
    return ReportResult(content, filename, CONTENT_TYPES[fmt], row_count, report_type, fmt, _clean_filters(filters))


def record_run(session: Session, result: ReportResult, triggered_by: str, file_path: str | None = None) -> ReportRun:
    """Add the report_runs row (not committed)."""
    run = ReportRun(report_type=result.report_type, file_path=file_path or result.filename, row_count=result.row_count,
                    filters_json={**result.filters, "format": result.fmt}, triggered_by=triggered_by)
    session.add(run)
    session.flush()
    return run


def generate_report(engine: Engine, report_type: str, filters: dict[str, Any], fmt: str, triggered_by: str,
                    save_dir: str | Path | None = None) -> ReportResult:
    """Build, optionally save to ``save_dir`` and log one report; commits. Runs from a plain script (no web server)."""
    with Session(engine) as session:
        result = build_report(session, report_type, filters, fmt, triggered_by)
        path = None
        if save_dir:
            folder = Path(save_dir)
            folder.mkdir(parents=True, exist_ok=True)
            path = folder / result.filename
            path.write_bytes(result.content)
        record_run(session, result, triggered_by, str(path) if path else None)
        session.commit()
    return result
