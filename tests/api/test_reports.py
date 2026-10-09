"""Exports: CSV/Excel round trips, role rules, the row cap, report_runs, EXPORT audit rows, agreement with analytics."""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine, text
from sqlalchemy.orm import Session

from app.services import audit
from etl import reports
from tests.api.conftest import auth_header
from tests.api.helpers import admission_of, load_forty

WINDOW = {"date_from": "2005-01-01", "date_to": "2005-06-30"}
AGGREGATED = ("admissions_trend", "readmission_summary", "length_of_stay_summary")


@pytest.fixture()
def api(client: TestClient, users_db: Engine) -> TestClient:
    with Session(users_db) as s, s.begin():
        load_forty(s)
    return client


def export(client: TestClient, user: str, report: str, fmt: str = "csv", **params):
    return client.get("/reports/export", params={"report": report, "format": fmt, **params}, headers=auth_header(client, user))


def as_csv(response) -> pd.DataFrame:
    return pd.read_csv(io.StringIO(response.content.decode("utf-8")), dtype={"diag_1": str, "diag_2": str, "diag_3": str})


def count(db: Engine, table: str, where: str = "1=1") -> int:
    with db.connect() as conn:
        return conn.execute(text(f"SELECT COUNT(*) FROM {table} WHERE {where}")).scalar_one()


# ---- downloads --------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("report", reports.REPORT_TYPES)
@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_every_report_downloads_in_both_formats(api: TestClient, report: str, fmt: str) -> None:
    r = export(api, "ops", report, fmt)
    assert r.status_code == 200 and len(r.content) > 100
    assert r.headers["content-type"].startswith(reports.CONTENT_TYPES[fmt])
    name = re.search(r'filename="([^"]+)"', r.headers["content-disposition"]).group(1)
    assert re.fullmatch(rf"{report}_\d{{8}}_\d{{6}}\.{fmt}", name) and " " not in name
    assert int(r.headers["x-row-count"]) > 0 and r.headers["x-request-id"]


def test_csv_is_clean_header_plus_data_only(api: TestClient) -> None:
    r = export(api, "admin", "readmission_summary", "csv")
    first = r.content.decode().splitlines()[0]
    assert first == "breakdown,label,encounters,eligible_encounters,readmitted_30d,readmission_rate,small_n,avg_length_of_stay"
    assert not any(line.startswith(("#", "generated", "Admission dates")) for line in r.content.decode().splitlines())
    df = as_csv(r)
    assert len(df) == int(r.headers["x-row-count"]) and set(df["breakdown"]) == {
        "age_group", "admission_type", "admission_source", "discharge_disposition", "specialty", "month"}
    assert df["encounters"].dtype.kind == "i" and df["readmission_rate"].dtype.kind == "f"  # numbers are numbers


def test_totals_equal_the_analytics_endpoints(api: TestClient) -> None:
    h = auth_header(api, "ana")
    summary = api.get("/analytics/summary", params=WINDOW, headers=h).json()["data"]
    df = as_csv(export(api, "ana", "readmission_summary", "csv", **WINDOW))
    for breakdown, part in df.groupby("breakdown"):
        assert part["encounters"].sum() == summary["total_encounters"], breakdown
        assert part["eligible_encounters"].sum() == summary["eligible_encounters"], breakdown
    trend = as_csv(export(api, "ana", "admissions_trend", "csv", **WINDOW))
    assert trend[trend["granularity"] == "month"]["encounters"].sum() == summary["total_encounters"]
    assert trend[trend["granularity"] == "year"]["encounters"].sum() == summary["total_encounters"]
    los = as_csv(export(api, "ana", "length_of_stay_summary", "csv", **WINDOW))
    assert los[los["breakdown"] == "age_group"]["encounters"].sum() == summary["total_encounters"]
    assert (los[[f"los_{d}_days" for d in range(1, 15)]].sum(axis=1) == los["encounters"]).all()  # histogram adds up


def test_excel_workbook_structure_and_values(api: TestClient) -> None:
    r = export(api, "ops", "readmission_summary", "xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "age_group", "admission_type", "admission_source", "discharge_disposition", "specialty",
                             "month", "Metadata"]
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=2)}
    assert summary["Total encounters"] == 39 and summary["Unique patients"] == 37 and summary["Eligible encounters (readmission denominator)"] == 36
    assert summary["30-day readmission rate (eligible)"] == pytest.approx(5 / 36, abs=1e-6)
    meta = {row[0].value: row[1].value for row in wb["Metadata"].iter_rows(min_row=2)}
    assert "SIMULATED" in meta["dates"] and "eligible encounters" in meta["denominator"] and meta["triggered_by"] == "ops"
    assert meta["row_count"] == int(r.headers["x-row-count"]) and meta["filters"] == "none"
    assert datetime.fromisoformat(meta["generated_at"])
    ages = wb["age_group"]
    assert ages["A1"].font.bold and ages.freeze_panes == "A2" and ages.column_dimensions["A"].width >= 12  # styled header, widths set
    assert isinstance(ages["B2"].value, int) and isinstance(ages["E2"].value, float)  # numbers, not text
    assert ages["E2"].number_format == "0.00%"


def test_excel_metadata_records_the_filters_and_dates_are_real_dates(api: TestClient) -> None:
    r = export(api, "admin", "encounters", "xlsx", **WINDOW)
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Encounters", "Metadata"]
    meta = {row[0].value: row[1].value for row in wb["Metadata"].iter_rows(min_row=2)}
    assert meta["filters"] == "date_from=2005-01-01, date_to=2005-06-30"
    sheet = wb["Encounters"]
    header = [c.value for c in sheet[1]]
    date_cell = sheet.cell(row=2, column=header.index("admission_date") + 1)
    assert isinstance(date_cell.value, datetime) and date_cell.number_format == "yyyy-mm-dd"  # a date, not a string
    assert isinstance(sheet.cell(row=2, column=header.index("time_in_hospital") + 1).value, int)


# ---- the encounter-level export: privacy, counts, caps ------------------------------------------------------------
def test_encounter_export_matches_the_list_api_and_exposes_no_demographics(api: TestClient) -> None:
    listing = api.get("/encounters", params={**WINDOW, "page_size": 100}, headers=auth_header(api, "admin")).json()
    for fmt in ("csv", "xlsx"):
        r = export(api, "admin", "encounters", fmt, **WINDOW)
        df = as_csv(r) if fmt == "csv" else pd.read_excel(io.BytesIO(r.content), sheet_name="Encounters")
        assert len(df) == listing["total"] == int(r.headers["x-row-count"])
        assert sorted(df["encounter_id"]) == sorted(i["encounter_id"] for i in listing["items"])
        assert list(df.columns) == reports.ENCOUNTER_EXPORT_COLUMNS
        assert not {"race", "gender"} & set(df.columns)
    raw = export(api, "admin", "encounters", "csv").content.decode().lower()
    assert "race" not in raw and "gender" not in raw and "caucasian" not in raw and "female" not in raw  # no value either
    full = as_csv(export(api, "admin", "encounters", "csv"))
    assert len(full) == 39 and full[full["encounter_id"] == 1]["diag_1"].iloc[0] == "250.00"  # ICD codes stay text (not 250.0)


def test_encounter_export_columns_are_an_explicit_allow_list() -> None:
    assert not reports.FORBIDDEN_COLUMNS & set(reports.ENCOUNTER_EXPORT_COLUMNS)
    assert "source_batch_id" not in reports.ENCOUNTER_EXPORT_COLUMNS
    source = (Path(reports.__file__)).read_text(encoding="utf-8")
    assert "fastapi" not in source.lower().replace("web-framework", "")  # callable without any web types


def test_filters_apply_to_the_encounter_export(api: TestClient) -> None:
    only = as_csv(export(api, "ops", "encounters", "csv", admission_type_id=2))
    assert len(only) == 9 and set(only["admission_type"]) == {"Urgent"}
    old = as_csv(export(api, "ops", "encounters", "csv", age_group="[80-90)"))
    assert len(old) == 14 and set(old["age_group"]) == {"[80-90)"}


def test_row_cap_returns_413_with_a_helpful_message(api: TestClient, monkeypatch: pytest.MonkeyPatch, users_db: Engine) -> None:
    monkeypatch.setattr(reports, "MAX_ENCOUNTER_EXPORT_ROWS", 10)
    r = export(api, "admin", "encounters", "csv")
    assert r.status_code == 413 and r.json()["error"]["code"] == "export_too_large"
    assert "limit is 10" in r.json()["error"]["message"] and "date_from" in r.json()["error"]["message"]
    assert r.json()["error"]["details"] == {"rows": 39, "limit": 10}
    assert count(users_db, "report_runs") == 0 and count(users_db, "audit_logs", "action = 'EXPORT'") == 0  # nothing recorded
    assert export(api, "admin", "encounters", "csv", admission_type_id=2).status_code == 200  # a narrower export (9 rows) passes
    assert export(api, "admin", "readmission_summary", "csv").status_code == 200  # the cap only applies to encounter-level files


# ---- roles ----------------------------------------------------------------------------------------------------------
def test_role_rules(api: TestClient) -> None:
    for fmt in ("csv", "xlsx"):
        assert export(api, "ana", "encounters", fmt).status_code == 403
        for user in ("admin", "ops"):
            assert export(api, user, "encounters", fmt).status_code == 200
        for report in AGGREGATED:
            for user in ("admin", "ops", "ana"):
                assert export(api, user, report, fmt).status_code == 200
    denied = export(api, "ana", "encounters")
    assert set(denied.json()["error"]) == {"code", "message", "details"} and "patient_nbr" not in denied.text
    assert api.get("/reports/export", params={"report": "encounters", "format": "csv"}).status_code == 401


@pytest.mark.parametrize("params", [{"report": "patients"}, {"report": "encounters", "format": "pdf"}, {"format": "csv"}, {}])
def test_invalid_export_parameters_are_422(api: TestClient, params: dict) -> None:
    r = api.get("/reports/export", params=params, headers=auth_header(api, "admin"))
    assert r.status_code == 422 and set(r.json()["error"]) == {"code", "message", "details"}


def test_history_is_for_administrator_and_clinical_ops(api: TestClient) -> None:
    export(api, "ana", "admissions_trend", "csv")
    assert api.get("/reports/history", headers=auth_header(api, "ana")).status_code == 403
    assert api.get("/reports/history").status_code == 401
    for user in ("admin", "ops"):
        assert api.get("/reports/history", headers=auth_header(api, user)).status_code == 200


# ---- bookkeeping: report_runs and audit -------------------------------------------------------------------------------
def test_every_export_writes_a_report_run_and_only_encounter_exports_write_an_audit_row(api: TestClient, users_db: Engine) -> None:
    export(api, "ana", "admissions_trend", "xlsx", **WINDOW)
    export(api, "ops", "readmission_summary", "csv")
    assert count(users_db, "report_runs") == 2 and count(users_db, "audit_logs", "action = 'EXPORT'") == 0
    r = export(api, "ops", "encounters", "csv", **WINDOW)
    assert count(users_db, "report_runs") == 3 and count(users_db, "audit_logs", "action = 'EXPORT'") == 1
    with users_db.connect() as conn:
        run = conn.execute(text("SELECT report_type, file_path, row_count, filters_json, triggered_by FROM report_runs ORDER BY report_id DESC LIMIT 1")).one()
        trail = conn.execute(text("SELECT username, role, entity_type, entity_id, after_json, request_id FROM audit_logs WHERE action = 'EXPORT'")).one()
    assert (run[0], run[2], run[4]) == ("encounters", int(r.headers["x-row-count"]), "ops") and run[1].startswith("encounters_")
    assert "2005-01-01" in str(run[3]) and "csv" in str(run[3])
    assert (trail[0], trail[1], trail[2]) == ("ops", "clinical_ops", "report") and trail[3] == run[1]
    assert trail[5] == r.headers["x-request-id"] and f'"rows": {run[2]}' in str(trail[4]).replace("'", '"')


def test_history_lists_what_was_exported(api: TestClient) -> None:
    export(api, "ops", "admissions_trend", "csv")
    export(api, "admin", "encounters", "csv", **WINDOW)
    body = api.get("/reports/history", headers=auth_header(api, "admin")).json()
    assert body["total"] == 2 and [i["report_type"] for i in body["items"]] == ["encounters", "admissions_trend"]
    only = api.get("/reports/history", params={"report_type": "encounters"}, headers=auth_header(api, "admin")).json()
    assert only["total"] == 1 and only["items"][0]["triggered_by"] == "admin"
    assert api.get("/reports/history", params={"date_from": "2999-01-01"}, headers=auth_header(api, "admin")).json()["total"] == 0


def test_a_failing_audit_write_leaves_no_report_run(api: TestClient, users_db: Engine, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(audit, "record", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("audit down")))
    admin = auth_header(api, "admin")
    r = api.get("/reports/export", params={"report": "encounters", "format": "csv"}, headers=admin)
    assert r.status_code == 500 and count(users_db, "report_runs") == 0  # no export is logged without its audit row


def test_soft_deleted_encounters_are_not_exported(api: TestClient) -> None:
    before = len(as_csv(export(api, "admin", "encounters", "csv")))
    api.delete("/encounters/1", headers=auth_header(api, "admin"))
    after = as_csv(export(api, "admin", "encounters", "csv"))
    assert len(after) == before - 1 and 1 not in set(after["encounter_id"]) and 40 not in set(after["encounter_id"])


# ---- the generator without HTTP ---------------------------------------------------------------------------------------
def test_generator_runs_from_a_plain_script_and_logs_the_run(users_db: Engine, tmp_path: Path) -> None:
    with Session(users_db) as s, s.begin():
        load_forty(s)
    result = reports.generate_report(users_db, "readmission_summary", {"date_to": date(2005, 12, 31)}, "xlsx", "airflow", save_dir=tmp_path)
    saved = tmp_path / result.filename
    assert saved.exists() and saved.read_bytes() == result.content and result.filename.endswith(".xlsx")
    with users_db.connect() as conn:
        run = conn.execute(text("SELECT triggered_by, file_path, row_count FROM report_runs")).one()
    assert run[0] == "airflow" and run[1] == str(saved) and run[2] == result.row_count
    assert count(users_db, "audit_logs") == 0  # no user, no audit row for a scheduled aggregated report
    with pytest.raises(ValueError):
        reports.generate_report(users_db, "nonsense", {}, "csv", "airflow")
